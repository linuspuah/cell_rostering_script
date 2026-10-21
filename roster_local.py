"""
Groop Groop cell roster generator - offline Excel version.

Same rules as roster.py, but reads the form responses and the committee
schedule from a local .xlsx file (e.g. the Google Sheet downloaded via
File > Download > Microsoft Excel) and writes the roster to a new .xlsx file.
No Google account or internet connection needed.

Keep this file in the same folder as roster.py - the rostering rules live there.

Usage
  python roster_local.py           # input/responses.xlsx -> output/roster_output.xlsx
  python roster_local.py "Groop Groop Responses.xlsx"
  python roster_local.py responses.xlsx --schedule-file schedule.xlsx
  python roster_local.py responses.xlsx --output roster_2027.xlsx
"""

import argparse
import datetime as dt
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from roster import (RESPONSES_TAB, ROSTER_TAB, SCHEDULE_TAB, SUMMARY_TAB, build_roster,
                    parse_members, parse_schedule, roster_table, summary_table)

HERE = Path(__file__).resolve().parent
DEFAULT_INPUT = HERE / "input" / "responses.xlsx"
DEFAULT_OUTPUT = HERE / "output" / "roster_output.xlsx"


def cell_str(v) -> str:
    """Excel gives back real dates/numbers; turn them into the text roster.py expects."""
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_tab(path: Path, tab: str) -> list[list[str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if tab not in wb.sheetnames:
        sys.exit(f"'{path.name}' has no tab named '{tab}'. Tabs found: {wb.sheetnames}")
    rows = [[cell_str(v) for v in row] for row in wb[tab].iter_rows(values_only=True)]
    wb.close()
    rows = [r for r in rows if any(r)]  # drop fully blank rows
    if not rows:
        sys.exit(f"Tab '{tab}' in '{path.name}' is empty.")
    return rows


def write_workbook(path: Path, tabs: dict[str, list[list]]):
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in tabs.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(width + 2, 60)
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description="Generate the Groop Groop cell roster from a local Excel file.")
    ap.add_argument("responses_file", type=Path, nargs="?", default=DEFAULT_INPUT,
                    help=f"Excel copy of the form responses (default: {DEFAULT_INPUT})")
    ap.add_argument("--responses-tab", default=RESPONSES_TAB)
    ap.add_argument("--schedule-file", type=Path,
                    help="Excel file with the committee schedule (default: same file as responses)")
    ap.add_argument("--schedule-tab", default=SCHEDULE_TAB)
    ap.add_argument("--output", type=Path, default=DEFAULT_OUTPUT,
                    help=f"where to save the roster (default: {DEFAULT_OUTPUT})")
    args = ap.parse_args()

    if not args.responses_file.exists():
        sys.exit(f"Input file not found: {args.responses_file}\n"
                 f"Save the downloaded Google Sheet there, or pass its path as an argument.")
    args.output.parent.mkdir(parents=True, exist_ok=True)

    resp_rows = read_tab(args.responses_file, args.responses_tab)
    sched_rows = read_tab(args.schedule_file or args.responses_file, args.schedule_tab)

    members = parse_members(resp_rows)
    weeks = parse_schedule(sched_rows)
    print(f"{len(members)} members, {len(weeks)} weeks")

    roster, load, counts = build_roster(members, weeks)
    try:
        write_workbook(args.output, {
            ROSTER_TAB: roster_table(weeks, roster),
            SUMMARY_TAB: summary_table(members, load, counts, weeks),
        })
    except PermissionError:
        sys.exit(f"Could not save {args.output} - close it in Excel and run again.")
    unfilled = sum(1 for wk in roster.values() for s in wk.values() if s.person.startswith("⚠"))
    print(f"Wrote {args.output}" + (f"  ({unfilled} unfilled slot(s) - see Notes column)" if unfilled else ""))


if __name__ == "__main__":
    main()
